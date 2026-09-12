"""xlsx 與 xlsb 的統一唯讀介面。extract 模組只透過這裡讀檔。"""
from __future__ import annotations
from pathlib import Path
from typing import Iterator
import warnings
import openpyxl

warnings.filterwarnings("ignore", message="Data Validation extension")


class Workbook:
    def __init__(self, sheetnames: list[str], reader, underlying=None):
        self.sheetnames = sheetnames
        self._reader = reader          # callable(sheet) -> Iterator[tuple]
        self._underlying = underlying  # openpyxl or pyxlsb workbook object

    def rows(self, sheet: str) -> Iterator[tuple]:
        return self._reader(sheet)

    def close(self) -> None:
        """Close underlying workbook resource. Idempotent and safe."""
        if self._underlying is not None and hasattr(self._underlying, 'close'):
            self._underlying.close()
            self._underlying = None

    def __enter__(self):
        return self

    def __exit__(self, *args):
        self.close()
        return False


def _open_xlsx(path: Path) -> Workbook:
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    return Workbook(list(wb.sheetnames), lambda s: wb[s].iter_rows(values_only=True), underlying=wb)


def _open_xlsb(path: Path):
    from pyxlsb import open_workbook as _pyxlsb_open
    return _pyxlsb_open(str(path))


def _xlsb_workbook(path: Path) -> Workbook:
    wb = _open_xlsb(path)
    names = list(wb.sheets)

    def reader(sheet: str) -> Iterator[tuple]:
        with wb.get_sheet(sheet) as sh:
            for row in sh.rows():
                yield tuple(c.v for c in row)
    return Workbook(names, reader, underlying=wb)


def open_workbook(path: str | Path) -> Workbook:
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".xlsx":
        return _open_xlsx(path)
    if ext == ".xlsb":
        return _xlsb_workbook(path)
    raise ValueError(f"unsupported workbook type: {path.name}")
